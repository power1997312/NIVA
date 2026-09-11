# -*- coding: utf-8 -*-
"""用例集渲染（TCG 第⑤步：交付文档）。

当前支持 **xlsx**（openpyxl 已在依赖内）。
docx 未实现：环境无 python-docx，且刻意不为此新增依赖 —— 显式返回降级原因，
不假装能出 Word。需要 docx 时再引入依赖（属 P5 打磨项）。
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Optional

from ... import config as C
from ...kernel.registry import ToolResult, degraded, fail, ok

HEADERS = [
    ("用例号", 14), ("页", 6), ("单元", 24), ("类型", 10), ("标签", 22),
    ("模板", 22), ("匹配", 10), ("目的", 40), ("前置条件", 34),
    ("步骤", 52), ("期望", 34), ("验收准则", 40), ("覆盖", 18),
    ("依据", 26), ("说明来源", 12), ("警告", 40),
]


def render_case_set_xlsx(case_set: dict, output_path: str | Path) -> ToolResult:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"

    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(name="宋体", size=11, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")

    for i, (h, _w) in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(vertical="center")
    for i, (_h, w) in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    def _steps(c: dict) -> str:
        return "\n".join(f"{s.get('no')}. {s.get('action')}\n   期望：{s.get('expected')}"
                         for s in c.get("steps") or [])

    def _vars(c: dict) -> str:
        vs = c.get("variables") or []
        if not vs:
            return ""
        return "；".join(
            f"{v.get('label')}（{v.get('cls')}，定值={v.get('setpoint')}"
            f"{'' if v.get('setpoint') is not None else '，待补'}）" for v in vs[:4])

    row = 2
    for c in case_set.get("cases") or []:
        vals = [
            c.get("case_id"), c.get("page"), c.get("unit_id"), c.get("kind"),
            c.get("label"),
            f"{(c.get('template') or {}).get('case_id')}"
            f"({'探针' if c.get('probe') else (c.get('template') or {}).get('match_kind')})",
            (c.get("template") or {}).get("match_kind"),
            c.get("principle"), c.get("preconditions"),
            _steps(c), _vars(c) or "；".join(str(x) for x in c.get("expected_values") or []),
            c.get("acceptance_criteria"),
            "、".join(c.get("coverage") or []),
            c.get("standard_ref"),
            c.get("narrative_source"),
            "；".join(c.get("warnings") or [])[:400],
        ]
        for i, v in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.alignment = wrap
            cell.font = Font(name="宋体", size=10)
        # 探针/降级行高亮：降级事实必须显式可见
        if c.get("probe") or c.get("narrative_source") == "template_fallback":
            ws.cell(row=row, column=16).fill = PatternFill("solid", fgColor="FFF2CC")
        row += 1

    # 汇总 sheet
    st = case_set.get("stats") or {}
    ws2 = wb.create_sheet("生成统计")
    ws2.append(["指标", "值"])
    for k, v in (st or {}).items():
        ws2.append([k, v])
    ws2.append(["跳过单元数", len(case_set.get("skipped") or [])])
    for w in case_set.get("warnings") or []:
        ws2.append(["警告", w])
    for i, w in enumerate([60, 90], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    wb.save(str(out))
    return ok({"xlsx_path": str(out), "bytes": out.stat().st_size,
               "cases": len(case_set.get("cases") or []),
               "stats": st})


def render_case_set(case_set: dict, output_path: str | Path,
                    fmt: str = "xlsx") -> ToolResult:
    if fmt == "xlsx":
        return render_case_set_xlsx(case_set, output_path)
    return degraded("NOT_IMPLEMENTED",
                    f"暂不支持 {fmt} 渲染：环境无 python-docx，为避免引入未评估依赖而显式降级；"
                    "当前可用格式为 xlsx",
                    data={"available_formats": ["xlsx"]})
